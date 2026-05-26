import streamlit as st
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font
from deep_translator import GoogleTranslator
import io
import time

st.set_page_config(page_title="Converter", page_icon="🚌")

st.markdown("""
<style>
    .stApp { background: linear-gradient(to right, #eef2ff, #f8fafc) !important; }
    .main-title { font-size: 48px !important; font-weight: 700 !important; color: #4338CA !important; text-align: center; margin-bottom: 20px; }
    .glass-card { background: white !important; padding: 30px; border-radius: 20px; box-shadow: 0 8px 32px rgba(31,38,135,0.15); }
</style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-title">🚌 Filllist To Customer Sheet</div>', unsafe_allow_html=True)
st.info("Filllist To Customer Sheet...")

if 'done' not in st.session_state:
    st.session_state.done = False
if 'output_data' not in st.session_state:
    st.session_state.output_data = None

uploaded_file = st.file_uploader("Upload filllist.xls or .xlsx", type=['xls', 'xlsx'])

if uploaded_file is None:
    st.session_state.done = False
    st.session_state.output_data = None

if uploaded_file:
    if not st.session_state.done:
        if st.button("🚀 Start Processing"):
            try:
                all_sheets = pd.read_excel(uploaded_file, sheet_name=None, header=None)
                translator = GoogleTranslator(source='en', target='ta')
                output = io.BytesIO()
                wb = Workbook()
                wb.remove(wb.active)

                thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                
                sheets_list = list(all_sheets.keys())
                total_sheets = len(sheets_list)
                my_bar = st.progress(0, text="Processing please wait...")

                for idx, route_no in enumerate(sheets_list):
                    df = all_sheets[route_no]
                    if df.empty: continue
                    
                    ws = wb.create_sheet(title=str(route_no))
                    bus_no = str(df.iloc[0, 1]) if df.shape[1] > 1 else "N/A"
                    locations = df.iloc[1:, 0].dropna().astype(str).tolist()

                    ws.merge_cells('A1:A2')
                    ws['A1'] = f"Route No: {route_no}"
                    ws['A1'].font = Font(bold=True); ws['A1'].alignment = Alignment(horizontal='left', vertical='center')
                    ws['A1'].border = thin_border; ws['A2'].border = thin_border 
                    
                    ws.merge_cells('B1:C2')
                    for r in range(1, 3):
                        for c in range(2, 4): ws.cell(row=r, column=c).border = thin_border
                    
                    ws.merge_cells('B3:C3')
                    ws['B3'] = f"Bus No: {bus_no}"; ws['B3'].font = Font(bold=True)
                    ws['B3'].alignment = Alignment(horizontal='center', vertical='center')
                    for c in range(1, 4): ws.cell(row=3, column=c).border = thin_border
                    ws['A3'].border = thin_border

                    labels = [("No of Singles", 1), ("10", 2), ("No of Bus-1", 3)]
                    for text, r_idx in labels:
                        cell = ws.cell(row=r_idx, column=4, value=text)
                        cell.border = thin_border; cell.alignment = Alignment(horizontal='center', vertical='center')

                    headers = ['Slot', 'Ad Location (Tamil)', 'Ad Location (English)', '']
                    for c_idx, text in enumerate(headers, 1):
                        cell = ws.cell(row=4, column=c_idx, value=text)
                        cell.font = Font(bold=True); cell.border = thin_border; cell.alignment = Alignment(horizontal='center')

                    for i, eng in enumerate(locations):
                        curr_row = 5 + i
                        try: tam = translator.translate(eng)
                        except: tam = eng
                        ws.cell(row=curr_row, column=1, value=f"A_{i+1}").border = thin_border
                        ws.cell(row=curr_row, column=2, value=tam).border = thin_border
                        ws.cell(row=curr_row, column=3, value=eng).border = thin_border
                        ws.cell(row=curr_row, column=4).border = thin_border

                    ws.column_dimensions['A'].width = 15
                    ws.column_dimensions['B'].width = 40
                    ws.column_dimensions['C'].width = 40
                    ws.column_dimensions['D'].width = 15
                    
                    my_bar.progress((idx + 1) / total_sheets)

                wb.save(output)
                st.session_state.output_data = output.getvalue()
                st.session_state.done = True
                st.rerun()

            except Exception as e:
                st.error(f"Error: {e}")

    if st.session_state.done:
        st.success("✅ Done!")
        st.download_button(
            label="Download Final Excel",
            data=st.session_state.output_data,
            file_name="Customer_Sheet_Final.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        if st.button("🔄 Process Another File"):
            st.session_state.done = False
            st.rerun()